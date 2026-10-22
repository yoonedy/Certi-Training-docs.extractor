import streamlit as st
import os
from pptx import Presentation
from docx import Document
import io

st.title("🎓 교육 행정 자동화 시스템 (GEMS)")
st.write("회사 양식(Template) 기반 수료증 및 교육확인서 자동 생성기")

# 아웃룩 텍스트 입력창
email_text = st.text_area("아웃룩 이메일 본문을 여기에 붙여넣으세요:", height=200, 
placeholder="""한국지역난방공사 충청지사 ㅇㅇ사업부
홍길동 과장
홍길순 대리
날짜 : 2026.10.22
강사명 : 김강사
교육명 : DCS 제어시스템 기초 교육
교육장 정보 : 본사 3층 대회의실""")

if st.button("파일 생성하기"):
    if not email_text.strip():
        st.warning("이메일 내용을 입력해 주세요!")
    else:
        # 아웃룩 텍스트 파싱
        lines = email_text.strip().split("\n")
        
        date_str = "2026.10.22"
        instructor = "김강사"
        course_name = "교육과정"
        location = "본사"
        company_info = "미정"
        students = []

        for line in lines:
            if "날짜" in line:
                date_str = line.split(":")[-1].strip()
            elif "강사명" in line:
                instructor = line.split(":")[-1].strip()
            elif "교육명" in line:
                course_name = line.split(":")[-1].strip()
            elif "교육장" in line:
                location = line.split(":")[-1].strip()
            elif any(keyword in line for keyword in ["과장", "대리", "사원", "부장", "차장", "주임"]):
                parts = line.strip().split()
                if len(parts) >= 2:
                    students.append({"name": parts[0], "rank": parts[1]})
            elif not any(k in line for k in [":", "날짜", "강사명", "교육명", "교육장"]):
                if not company_info or company_info == "미정":
                    company_info = line.strip()

        file_date_str = date_str.replace(".", "")

        # --- PPT 수료증 치환 로직 (개선됨) ---
        ppt_path = "template.pptx"
        if os.path.exists(ppt_path):
            prs = Presentation(ppt_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            # 단락 전체 텍스트에서 치환자 변경
                            full_text = paragraph.text
                            if any(k in full_text for k in ["{{course_name}}", "{{date}}", "{{instructor}}"]):
                                full_text = full_text.replace("{{course_name}}", course_name)
                                full_text = full_text.replace("{{date}}", date_str)
                                full_text = full_text.replace("{{instructor}}", instructor)
                                # 첫 번째 런에 새 텍스트를 넣고 나머지는 비움
                                if len(paragraph.runs) > 0:
                                    paragraph.runs[0].text = full_text
                                    for r in paragraph.runs[1:]:
                                        r.text = ""
            
            ppt_io = io.BytesIO()
            prs.save(ppt_io)
            ppt_io.seek(0)
            ppt_available = True
        else:
            ppt_available = False

        # --- Word 교육확인서 치환 로직 (개선됨) ---
        doc_path = "template.docx"
        if os.path.exists(doc_path):
            doc = Document(doc_path)
            
            for p in doc.paragraphs:
                if any(k in p.text for k in ["{{course_name}}", "{{date}}", "{{instructor}}", "{{location}}"]):
                    p.text = p.text.replace("{{course_name}}", course_name)
                    p.text = p.text.replace("{{date}}", date_str)
                    p.text = p.text.replace("{{instructor}}", instructor)
                    p.text = p.text.replace("{{location}}", location)

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if "{{company}}" in cell.text:
                            cell.text = cell.text.replace("{{company}}", company_info)

            doc_io = io.BytesIO()
            doc.save(doc_io)
            doc_io.seek(0)
            doc_available = True
        else:
            doc_available = False

        st.success("파일 생성 완료!")

        if ppt_available:
            st.download_button(
                label="📥 수료증 다운로드 (PPTX)",
                data=ppt_io,
                file_name=f"{course_name}_수료증_{file_date_str}.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
            )
        else:
            st.warning("⚠️ 'GEMS' 폴더에 `template.pptx` 파일이 없습니다.")

        if doc_available:
            st.download_button(
                label="📥 교육확인서 다운로드 (DOCX)",
                data=doc_io,
                file_name=f"{course_name}_교육확인서_{file_date_str}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        else:
            st.warning("⚠️ 'GEMS' 폴더에 `template.docx` 파일이 없습니다.")