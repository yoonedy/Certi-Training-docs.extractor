"""ABB KR AUEN TV 대시보드 서버.

    python server.py            → http://localhost:8080 (TV 브라우저 전체화면으로 열기)

- /api/schedule : 스케줄 엑셀(.xlsb/.xlsx) → JSON. 파일이 바뀌면 자동으로 다시 읽음
- /api/today    : 오늘 음력 날짜 · 24절기 (calendar_kr.py)
- /api/memo     : 지도 하단 메모 (memo.file 텍스트 파일 → 없으면 config의 memo.text)
- /api/notices  : Power Automate가 notices_dir에 저장한 전체공지 JSON 목록
- /api/menu     : menu_weekly_dir 의 최신 주간 식단표 이미지 + 글자 인식 결과(text)
- /api/news     : ABB 소식 / 경제뉴스 RSS (15분 캐시)
- /api/weather  : Open-Meteo 현재 기온 (10분 캐시)
- /api/stocks   : Yahoo Finance 시세 (5분 캐시)

엑셀·공지·식단표 폴더는 SharePoint 문서 라이브러리를 OneDrive로 동기화한 로컬 경로를 지정하면
SharePoint에 올리는 즉시 TV에 반영된다.
"""
import calendar
import datetime as dt
import email.utils
import html
import json
import mimetypes
import re
import sys
import threading
import time
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

ROOT = Path(__file__).resolve().parent
STATIC = ROOT / "static"
CONFIG_PATH = ROOT / "config.json"
EXCEL_EPOCH = dt.date(1899, 12, 30)
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}

# 엑셀 헤더 → JSON 키. 열 순서가 바뀌어도 헤더 이름으로 찾는다.
COLUMNS = {
    "프로젝트 번호": "no",
    "프로젝트 명": "name",
    "담당자(쉼표구분)": "members",
    "ABB Work 시작": "abbStart",
    "ABB Work 종료": "abbEnd",
    "FAT 시작": "fatStart",
    "FAT 종료": "fatEnd",
    "납품 시작": "delStart",
    "납품 종료": "delEnd",
    "시운전 시작": "siteStart",
    "시운전 종료": "siteEnd",
    "프로젝트 종료": "end",
    # 선택 열: 있으면 config의 키워드 매칭보다 우선
    "현장": "site",
    "표시명": "title",
}


def load_config():
    return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))


def resolve(path_str):
    p = Path(path_str)
    return p if p.is_absolute() else ROOT / p


# ---------------------------------------------------------------- 엑셀

def to_date(v):
    """엑셀 일련번호 / 문자열 날짜 → 'YYYY-MM-DD'. '2026/11/31'처럼 없는 날짜는 그 달 말일로 보정."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        if v < 20000:  # 날짜가 아닌 숫자
            return None
        return (EXCEL_EPOCH + dt.timedelta(days=int(v))).isoformat()
    s = str(v).strip()
    m = re.match(r"^(\d{4})[./-](\d{1,2})[./-](\d{1,2})", s)
    if m:
        y, mo, d = map(int, m.groups())
    else:
        m = re.match(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{4})", s)  # MM/DD/YYYY
        if not m:
            return None
        mo, d, y = map(int, m.groups())
    if not (1 <= mo <= 12):
        return None
    d = max(1, min(d, calendar.monthrange(y, mo)[1]))
    return dt.date(y, mo, d).isoformat()


def read_rows(path, sheet):
    if path.suffix.lower() == ".xlsb":
        from pyxlsb import open_workbook
        with open_workbook(str(path)) as wb:
            with wb.get_sheet(sheet) as sh:
                return [[c.v for c in row] for row in sh.rows()]
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        return [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()


def parse_schedule(path, sheet):
    rows = read_rows(path, sheet)
    header_idx, colmap = None, {}
    for i, row in enumerate(rows[:10]):
        names = [str(c).strip() if c is not None else "" for c in row]
        if "프로젝트 번호" in names:
            header_idx = i
            colmap = {j: COLUMNS[n] for j, n in enumerate(names) if n in COLUMNS}
            break
    if header_idx is None:
        raise ValueError(f"'{sheet}' 시트에서 '프로젝트 번호' 헤더를 찾지 못했습니다")

    projects, warnings = [], []
    for r, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        rec = {}
        for j, key in colmap.items():
            v = row[j] if j < len(row) else None
            if key in ("no", "name", "members", "site", "title"):
                rec[key] = str(v).strip() if v not in (None, "") else ""
            else:
                rec[key] = to_date(v)
                if v not in (None, "") and rec[key] is None:
                    warnings.append(f"{r}행 '{key}' 날짜 해석 실패: {v!r}")
                elif isinstance(v, str) and rec[key]:
                    warnings.append(f"{r}행 '{key}' 문자열 날짜 {v!r} → {rec[key]}")
        if not rec.get("no") and not rec.get("name"):
            continue
        rec["members"] = [m.strip() for m in rec.get("members", "").split(",") if m.strip()]
        rec["row"] = r
        projects.append(rec)
    return projects, warnings


_schedule_cache = {"mtime": None, "data": None}
_schedule_lock = threading.Lock()


def get_schedule(cfg):
    path = resolve(cfg["excel_path"])
    if not path.exists():
        return {"error": f"엑셀 파일 없음: {path}", "projects": []}
    mtime = path.stat().st_mtime
    with _schedule_lock:
        if _schedule_cache["mtime"] != mtime:
            try:
                projects, warnings = parse_schedule(path, cfg.get("excel_sheet", "입력양식"))
            except Exception as e:  # 저장 중(잠김)이면 이전 데이터를 유지
                if _schedule_cache["data"]:
                    return {**_schedule_cache["data"], "stale": str(e)}
                return {"error": str(e), "projects": []}
            _schedule_cache["mtime"] = mtime
            _schedule_cache["data"] = {
                "file": path.name,
                "updated": dt.datetime.fromtimestamp(mtime).isoformat(timespec="minutes"),
                "projects": projects,
                "warnings": warnings,
            }
        return _schedule_cache["data"]


# ---------------------------------------------------------------- 공지 / 식단표

NOTICE_HEADERS = {  # 엑셀 머리글 → 키 (Power Automate '표에 행 추가'에서 쓴 이름)
    "receiveddatetime": "received", "received": "received", "받은 시간": "received", "날짜": "received",
    "subject": "subject", "제목": "subject",
    "from": "from", "보낸 사람": "from", "sender": "from",
    "torecipients": "to", "to": "to", "받는 사람": "to", "ccrecipients": "cc", "cc": "cc", "참조": "cc",
    "body": "body", "bodypreview": "body", "본문": "body",
}
NOTICE_DEFAULT_COLS = ["received", "subject", "from", "to", "cc", "body"]  # 머리글이 없으면 A~F 순서로 봄


def _when(s):
    try:
        d = dt.datetime.fromisoformat(str(s).replace("Z", "+00:00"))
        return d if d.tzinfo else d.astimezone()
    except ValueError:
        return dt.datetime.min.replace(tzinfo=dt.timezone.utc)


def _notices_from_folder(folder):
    items = []
    for f in folder.glob("*.json") if folder.exists() else []:
        try:
            d = json.loads(f.read_text(encoding="utf-8-sig"))
        except Exception:
            continue
        for n in d if isinstance(d, list) else [d]:
            items.append({
                "subject": n.get("subject") or n.get("Subject") or f.stem,
                "from": n.get("from") or n.get("From") or "",
                "to": n.get("toRecipients") or n.get("to") or "",
                "cc": n.get("ccRecipients") or n.get("cc") or "",
                "received": n.get("received") or n.get("receivedDateTime") or
                            dt.datetime.fromtimestamp(f.stat().st_mtime).isoformat(),
                "body": n.get("body") or n.get("bodyPreview") or n.get("Body") or "",
            })
    return items


def _notices_from_excel(src):
    """Power Automate가 메일을 한 줄씩 쌓는 엑셀 (받은 시간 · 제목 · 보낸 사람 · [받는 사람] · [본문])"""
    path = resolve(src["path"])
    if not path.is_file():
        return []
    if path.suffix.lower() == ".xlsb":
        rows = read_rows(path, src.get("sheet") or 1)
    else:
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        try:
            ws = wb[src["sheet"]] if src.get("sheet") else wb.worksheets[0]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    if not rows:
        return []
    head = [NOTICE_HEADERS.get(str(c).strip().lower()) if c is not None else None for c in rows[0]]
    if "subject" in head:
        cols, rows = head, rows[1:]
    else:
        cols = NOTICE_DEFAULT_COLS
    items = []
    for r in rows:
        rec = {k: (str(r[i]).strip() if i < len(r) and r[i] is not None else "") for i, k in enumerate(cols) if k}
        if rec.get("subject"):
            items.append(rec)
    return items


def recipients(n):
    """받는 사람 + 참조 주소 목록 (Power Automate는 ';'로 이어 붙여 줌), 중복 제거"""
    found = re.findall(r"[\w.+'-]+@[\w-]+(?:\.[\w-]+)+", f'{n.get("to", "")};{n.get("cc", "")}')
    return sorted({a.lower() for a in found})


def is_notice(n, rules):
    """진짜 전체공지인지:
    (1) 받는 사람 조건 — require_to 주소(전 직원 배포 목록)가 있거나, 받는 사람+참조가 min_recipients명 이상
        (둘 다 설정하면 하나만 맞아도 통과. 받는 사람 정보가 아예 없는 메일엔 인원 조건을 적용하지 않음)
    (2) 자동 알림·회신·테스트 같은 보낸 사람/제목 패턴은 제외 (3) include_subject가 있으면 그중 하나는 맞아야 함"""
    find = lambda pats, text: any(re.search(p, text or "", re.I) for p in pats)
    need = rules.get("require_to") or []
    min_n = int(rules.get("min_recipients") or 0)
    rcpt = recipients(n)
    if need or (min_n and rcpt):
        to_list = need and find([re.escape(a) for a in need], ";".join(rcpt))
        many = min_n and len(rcpt) >= min_n
        if not (to_list or many):
            return False
    if find(rules.get("exclude_from", []), n.get("from", "")):
        return False
    if find(rules.get("exclude_subject", []), n.get("subject", "")):
        return False
    inc = rules.get("include_subject") or []
    return not inc or find(inc, n.get("subject", ""))


def get_notices(cfg, limit=8):
    items = _notices_from_folder(resolve(cfg["notices_dir"]))
    if (cfg.get("notices_excel") or {}).get("path"):
        items += _notices_from_excel(cfg["notices_excel"])
    rules = cfg.get("notices_filter", {})
    items = [n for n in items if is_notice(n, rules)]
    items.sort(key=lambda n: _when(n["received"]), reverse=True)
    out, seen = [], set()
    for n in items:  # 같은 공지를 두 번 보낸 경우(재발송) 최신 것만
        key = re.sub(r"\s+", "", n["subject"]).lower()
        if key in seen:
            continue
        seen.add(key)
        sender = n.get("from", "")
        out.append({"subject": n["subject"], "received": n["received"], "body": strip_html(n.get("body", ""))[:400],
                    "from": sender.split("@")[0] if "@" in sender else sender})
        if len(out) >= limit:
            break
    return {"items": out}


def latest_image(folder):
    if not folder.exists():
        return None
    files = [f for f in folder.rglob("*") if f.is_file() and f.suffix.lower() in IMAGE_EXT]
    if not files:
        return None
    f = max(files, key=lambda p: p.stat().st_mtime)
    return {"name": f.name, "rel": f.relative_to(folder).as_posix(), "updated": dt.datetime.fromtimestamp(f.stat().st_mtime).isoformat(timespec="minutes")}


def get_menu(cfg):
    weekly = latest_image(resolve(cfg["menu_weekly_dir"]))
    if weekly:
        weekly["url"] = "/files/menu-weekly/" + urllib.parse.quote(weekly["rel"])
    out = {"weekly": weekly}
    if weekly and cfg.get("menu_ocr", {}).get("enabled", True):
        out["text"] = menu_text(cfg, resolve(cfg["menu_weekly_dir"]) / weekly["rel"])
    return out


# 주간 식단표 글자 읽기는 몇 초 걸리므로 새 이미지가 보이면 뒤에서 한 번만 돌리고 결과를 기억해 둠
_ocr = {"key": None, "result": None, "running": False}
_ocr_lock = threading.Lock()


def menu_text(cfg, path):
    key = (str(path), path.stat().st_mtime)
    with _ocr_lock:
        if _ocr["key"] == key:
            return _ocr["result"]
        if _ocr["running"]:
            return None
        _ocr["running"] = True
    opts = cfg.get("menu_ocr", {})

    def work():
        try:
            import menu_ocr
            result = menu_ocr.read_menu(path, opts.get("tesseract_cmd", "tesseract"), opts.get("cutoff", 0.68))
        except Exception as e:  # Tesseract가 없거나 양식이 달라도 화면은 이미지로 계속 보여줌
            result = {"error": str(e)}
        with _ocr_lock:
            _ocr.update(key=key, result=result, running=False)

    threading.Thread(target=work, daemon=True).start()
    return None


# ---------------------------------------------------------------- 외부 데이터 (캐시)

_cache = {}


def cached(key, ttl, fn):
    hit = _cache.get(key)
    if hit and time.time() - hit[0] < ttl:
        return hit[1]
    try:
        val = fn()
        _cache[key] = (time.time(), val)
        return val
    except Exception as e:
        if hit:
            return {**hit[1], "stale": str(e)}
        return {"error": str(e)}


def http_get(url, timeout=10):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (TV dashboard)"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def strip_html(s):
    s = html.unescape(re.sub(r"<[^>]+>", " ", s or ""))  # &#x27; 같은 숫자 엔티티까지 글자로
    return re.sub(r"\s+", " ", s).strip()


def _rss_time(item):
    try:
        return email.utils.parsedate_to_datetime(item["date"]).timestamp()
    except (TypeError, ValueError):
        return 0


def fetch_rss(url, limit=8, newest_first=False):
    """newest_first: Google 뉴스 검색처럼 관련도 순인 피드를 최신순으로 다시 정렬"""
    root = ET.fromstring(http_get(url))
    items = []
    for it in root.iter("item"):
        title = strip_html(it.findtext("title") or "")
        source = it.findtext("source") or ""
        if source and title.endswith(" - " + source):  # Google News 제목 꼬리 제거
            title = title[: -len(source) - 3]
        items.append({"title": title, "source": source, "date": it.findtext("pubDate") or ""})
    if newest_first:
        items.sort(key=_rss_time, reverse=True)
    return {"items": items[:limit]}


def _anchor_title(tag, inner):
    """기사 링크(<a>) 안에서 제목만: 제목 태그(h1~h4·strong·*tit*) → 첫 글자 덩어리 → title 속성"""
    h = re.search(r'<(h[1-4]|strong)\b[^>]*>(.*?)</\1>', inner, re.S | re.I) or \
        re.search(r'<(\w+)\b[^>]*class="[^"]*(?:tit|ttl)[^"]*"[^>]*>(.*?)</\1>', inner, re.S | re.I)
    if h and len(strip_html(h.group(2))) >= 8:
        return strip_html(h.group(2))
    for chunk in re.split(r"<[^>]+>", inner):  # 썸네일·요약·날짜가 같이 들어 있어도 첫 덩어리가 제목
        if len(strip_html(chunk)) >= 8:
            return strip_html(chunk)
    alt = re.search(r'(?:title|aria-label)="([^"]{8,})"', tag)  # '자세히 보기' 같은 버튼이면 속성에서
    return strip_html(alt.group(1)) if alt else ""


def fetch_headlines(src, limit=8):
    """웹페이지에서 기사 링크 제목을 위에서부터 뽑는다 (ABB 코리아 홈페이지, 매일경제 경제 등).
    링크 주소가 link_regex(정규식)에 맞거나 link_pattern(기본 '/news/detail/')을 포함하는 <a>만 기사로 본다."""
    url = src["url"]
    regex = re.compile(src["link_regex"]) if src.get("link_regex") else None
    pattern = src.get("link_pattern", "/news/detail/")
    page = http_get(url).decode("utf-8", "replace")
    items, seen = [], set()
    for m in re.finditer(r'<a\b[^>]*href="([^"]+)"[^>]*>(.*?)</a>', page, re.S | re.I):
        href, inner = m.group(1), m.group(2)
        if not (regex.search(href) if regex else pattern in href):
            continue
        link = urllib.parse.urljoin(url, href)
        title = _anchor_title(m.group(0)[:m.group(0).find(">") + 1], inner)
        if not title or link in seen:  # 사진만 있는 링크는 건너뛰고 같은 기사의 제목 링크를 씀
            continue
        seen.add(link)
        items.append({"title": title, "source": src.get("source", "ABB 코리아"), "date": "", "link": link})
        if len(items) >= limit:
            break
    if not items:
        raise ValueError("기사 링크를 찾지 못함 (link_regex / link_pattern 확인)")
    return {"items": items}


def fetch_naver_section(src, limit=8):
    """네이버 뉴스 섹션(예: 경제 news.naver.com/section/101) 맨 위 '헤드라인' 묶음의 기사.
    기사 제목 링크는 class에 sa_text_title, 언론사는 sa_text_press. 헤드라인 묶음(as_headline)을
    못 찾으면 페이지의 기사 목록 앞부분을 쓴다."""
    page = http_get(src["url"]).decode("utf-8", "replace")
    m = re.search(r'class="[^"]*\bas_headline\b', page)
    if m:  # 헤드라인 묶음 ~ 다음 기사 묶음 전까지
        nxt = page.find('class="section_article', m.end())
        page = page[m.start(): nxt if nxt > 0 else None]
    items, seen = [], set()
    anchors = list(re.finditer(r'<a\b([^>]*)>(.*?)</a>', page, re.S | re.I))
    for k, a in enumerate(anchors):
        attrs = a.group(1)
        if "sa_text_title" not in attrs:
            continue
        href = re.search(r'href="([^"]+)"', attrs)
        title = strip_html(a.group(2))
        if not href or not title or href.group(1) in seen:
            continue
        seen.add(href.group(1))
        end = anchors[k + 1].start() if k + 1 < len(anchors) else len(page)
        press = re.search(r'class="[^"]*sa_text_press[^"]*"[^>]*>(.*?)<', page[a.end():end], re.S)
        items.append({"title": title, "source": strip_html(press.group(1)) if press else "네이버 뉴스",
                      "date": "", "link": href.group(1)})
        if len(items) >= limit:
            break
    if not items:
        raise ValueError("네이버 헤드라인 기사를 찾지 못함 (페이지 구조 변경?)")
    return {"items": items}


def fetch_news(src):
    """src: RSS 주소 문자열, 또는 {"type": "html" | "naver", "url", ..., "fallback_rss"}"""
    if isinstance(src, str):
        return fetch_rss(src)
    try:
        if src.get("type") == "naver":
            return fetch_naver_section(src)
        return fetch_headlines(src)
    except Exception:
        if src.get("fallback_rss"):
            return fetch_rss(src["fallback_rss"], newest_first=True)
        raise


WEATHER_CODES = {0: "맑음", 1: "대체로 맑음", 2: "구름 조금", 3: "흐림", 45: "안개", 48: "안개",
                 51: "이슬비", 53: "이슬비", 55: "이슬비", 61: "비", 63: "비", 65: "강한 비",
                 71: "눈", 73: "눈", 75: "강한 눈", 80: "소나기", 81: "소나기", 82: "강한 소나기",
                 95: "뇌우", 96: "뇌우", 99: "뇌우"}


def fetch_weather(w):
    q = urllib.parse.urlencode({
        "latitude": w["lat"], "longitude": w["lon"], "timezone": "Asia/Seoul",
        "current": "temperature_2m,weather_code,relative_humidity_2m",
        "daily": "temperature_2m_max,temperature_2m_min", "forecast_days": 1,
    })
    d = json.loads(http_get("https://api.open-meteo.com/v1/forecast?" + q))
    code = d["current"]["weather_code"]
    out = {
        "name": w.get("name", ""),
        "temp": d["current"]["temperature_2m"],
        "humidity": d["current"]["relative_humidity_2m"],
        "code": code,
        "text": WEATHER_CODES.get(code, ""),
        "max": d["daily"]["temperature_2m_max"][0],
        "min": d["daily"]["temperature_2m_min"][0],
    }
    try:  # 미세먼지는 실패해도 날씨는 그대로 보여줌
        out.update(fetch_air(w))
    except Exception as e:
        out["air_error"] = str(e)
    return out


# 환경부 통합대기 기준 (㎍/㎥): 좋음 / 보통 / 나쁨 / 매우나쁨
PM_GRADES = {"pm10": (30, 80, 150), "pm25": (15, 35, 75)}
GRADE_NAMES = ("좋음", "보통", "나쁨", "매우나쁨")


def pm_grade(kind, value):
    if value is None:
        return None
    return GRADE_NAMES[sum(value > limit for limit in PM_GRADES[kind])]


def fetch_air(w):
    q = urllib.parse.urlencode({"latitude": w["lat"], "longitude": w["lon"], "timezone": "Asia/Seoul",
                                "current": "pm10,pm2_5"})
    cur = json.loads(http_get("https://air-quality-api.open-meteo.com/v1/air-quality?" + q))["current"]
    pm10, pm25 = cur.get("pm10"), cur.get("pm2_5")
    g10, g25 = pm_grade("pm10", pm10), pm_grade("pm25", pm25)
    # 둘 중 나쁜 쪽을 대표 등급으로
    worst = max((g for g in (g10, g25) if g), key=GRADE_NAMES.index, default=None)
    return {"pm10": pm10, "pm25": pm25, "pm10_grade": g10, "pm25_grade": g25, "air_grade": worst}


def fetch_stocks(stocks):
    out = []
    for s in stocks:
        try:
            url = ("https://query1.finance.yahoo.com/v8/finance/chart/" + urllib.parse.quote(s["symbol"]) +
                   "?range=5d&interval=1d")
            meta = json.loads(http_get(url))["chart"]["result"][0]["meta"]
            price, prev = meta["regularMarketPrice"], meta.get("chartPreviousClose") or meta.get("previousClose")
            out.append({"name": s["name"], "symbol": s["symbol"], "price": price, "currency": meta.get("currency", ""),
                        "change": price - prev if prev else None,
                        "pct": (price - prev) / prev * 100 if prev else None})
        except Exception as e:
            out.append({"name": s["name"], "symbol": s["symbol"], "error": str(e)})
    return {"items": out}


# ---------------------------------------------------------------- HTTP

class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *a, **kw):
        super().__init__(*a, directory=str(STATIC), **kw)

    def log_message(self, fmt, *args):
        pass

    def send_json(self, obj, status=200):
        body = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def send_file(self, path):
        if not path.is_file():
            return self.send_error(404)
        data = path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", mimetypes.guess_type(path.name)[0] or "application/octet-stream")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        cfg = load_config()
        url = urllib.parse.urlparse(self.path)
        p = url.path
        if p == "/api/config":
            keys = ("rotation_seconds", "board_page_seconds", "refresh_minutes", "notices_refresh_hours", "imminent_days", "load_high",
                    "load_extreme", "sites", "aliases", "weather")
            return self.send_json({k: cfg.get(k) for k in keys})
        if p == "/api/schedule":
            return self.send_json(get_schedule(cfg))
        if p == "/api/today":
            import calendar_kr
            return self.send_json(cached("today:" + dt.date.today().isoformat(), 3600, calendar_kr.today_info))
        if p == "/api/memo":
            memo = dict(cfg.get("memo") or {})
            f = resolve(memo.pop("file", "") or "data/memo.txt")
            if f.is_file():
                memo["text"] = f.read_text(encoding="utf-8-sig").strip()
            return self.send_json(memo)
        if p == "/api/notices":
            return self.send_json(get_notices(cfg))
        if p == "/api/menu":
            return self.send_json(get_menu(cfg))
        if p == "/api/news":
            feed = urllib.parse.parse_qs(url.query).get("feed", [""])[0]
            src = cfg.get("news", {}).get(feed)
            if not src:
                return self.send_json({"error": "unknown feed"}, 404)
            return self.send_json(cached("news:" + feed, 900, lambda: fetch_news(src)))
        if p == "/api/weather":
            return self.send_json(cached("weather", 600, lambda: fetch_weather(cfg["weather"])))
        if p == "/api/stocks":
            return self.send_json(cached("stocks", 300, lambda: fetch_stocks(cfg.get("stocks", []))))
        if p.startswith("/files/menu-weekly/"):
            folder = resolve(cfg["menu_weekly_dir"])
            rel = urllib.parse.unquote(p.split("/", 3)[3])
            target = (folder / rel).resolve()
            if folder.resolve() not in target.parents:  # 경로 이탈 방지
                return self.send_error(404)
            return self.send_file(target)
        return super().do_GET()


def check_news(cfg):
    """python server.py --check-news : 뉴스 출처마다 실제로 읽히는지 확인 (대비용 RSS로 넘어갔는지도 표시)"""
    for name, src in cfg.get("news", {}).items():
        print(f"== {name}")
        try:
            if isinstance(src, dict):
                fn = fetch_naver_section if src.get("type") == "naver" else fetch_headlines
                items = fn(src)["items"]
                print("  원래 출처 OK:", src["url"])
            else:
                items = fetch_rss(src)["items"]
        except Exception as e:
            print("  원래 출처 실패:", e)
            items = []
            if isinstance(src, dict) and src.get("fallback_rss"):
                try:
                    items = fetch_rss(src["fallback_rss"], newest_first=True)["items"]
                    print("  → 대비용 RSS 사용:", src["fallback_rss"])
                except Exception as e2:
                    print("  → 대비용 RSS도 실패:", e2)
        for n in items[:8]:
            print(f"  - {n['title']}  ({n['source']})")


def main():
    cfg = load_config()
    if "--check-news" in sys.argv:
        return check_news(cfg)
    port = int(cfg.get("port", 8080))
    print(f"TV 대시보드: http://localhost:{port}")
    ThreadingHTTPServer(("0.0.0.0", port), Handler).serve_forever()


if __name__ == "__main__":
    main()
